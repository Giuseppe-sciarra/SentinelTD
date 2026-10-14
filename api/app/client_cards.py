"""Scheda cliente: lettura del foglio Excel/CSV dei domini e abbinamento ai provider.

Il foglio si riconosce dalle intestazioni, non dal nome o dall'ordine delle colonne:
- un foglio con "Dominio" e le colonne della posta (Email ospitate su, Come gestire, Pannello,
  Account, Webmail, DNS, Sito su, Note, e facoltativi Chi paga / Costo / Scadenza) = domini attivi;
- un foglio con "Dominio" e "Situazione" ma senza la colonna della posta = domini non attivi;
- un foglio con "Provider" e le procedure (Dove entrare, Crea, Password, Inoltri, Webmail) =
  istruzioni dei provider. Le righe di solo testo lungo di questo foglio diventano note del cliente.
Gli altri fogli (riepiloghi) si ignorano. Nessuna password: solo riferimenti.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime

EMPTY = {"", "—", "–", "-", "n/d", "nd", "n.d.", "none", "null"}

DOMAIN_FIELDS = [
    ("domain", ("dominio", "domain")),
    ("client", ("sottocliente", "cliente", "client")),
    ("mail_host", ("email ospitate", "email su", "posta su", "mail su", "email hosted", "mail host")),
    ("manage_howto", ("come gestire", "gestione caselle", "how to manage")),
    ("panel_url", ("link pannello", "pannello", "panel")),
    ("account_ref", ("account",)),
    ("webmail_url", ("webmail",)),
    ("dns_host", ("dns", "record mx")),
    ("site_host", ("sito su", "site on", "server sito", "sito")),
    ("status_note", ("situazione", "status")),
    ("renewal_payer", ("chi paga", "pagato da", "paid by")),
    ("renewal_cost", ("costo", "importo", "prezzo", "cost")),
    ("hosting_expires", ("scadenza", "expires", "rinnovo")),
    ("notes", ("note", "notes")),
]
PROVIDER_FIELDS = [
    ("name", ("provider", "fornitore")),
    ("login", ("dove entrare", "accesso", "login")),
    ("create_mailbox", ("crea", "elimina casella")),
    ("change_password", ("password",)),
    ("forwarding", ("inoltr", "risposta automatica", "forward")),
    ("webmail", ("webmail",)),
    ("notes", ("note", "notes")),
]
DOMAIN_TEXT = ("mail_host", "manage_howto", "panel_url", "account_ref", "webmail_url", "dns_host",
               "site_host", "notes", "status_note", "renewal_payer", "renewal_cost")
PROVIDER_TEXT = ("login", "create_mailbox", "change_password", "forwarding", "webmail", "notes")
# dove sta ogni colonna del foglio nel pannello (per riscrivere "colonna G del foglio ...")
FIELD_LABELS = {
    "domain": "Dominio", "mail_host": "Email ospitate su", "manage_howto": "Come gestire le caselle",
    "panel_url": "Link pannello", "account_ref": "Account da usare", "webmail_url": "Webmail",
    "dns_host": "DNS / record MX modificabili su", "site_host": "Sito su", "notes": "Note",
    "status_note": "Situazione", "renewal_payer": "Chi paga", "renewal_cost": "Costo", "hosting_expires": "Scadenza",
    "client": "Sottocliente",
}
PROVIDER_LABELS = {
    "name": "Nome", "login": "Dove entrare", "create_mailbox": "Crea / elimina casella", "change_password": "Cambia password",
    "forwarding": "Inoltri e risposta automatica", "webmail": "Webmail", "notes": "Note",
}
LIMITS = {"mail_host": 190, "account_ref": 255, "dns_host": 190, "site_host": 190, "renewal_payer": 120,
          "renewal_cost": 60, "name": 190}
_DOMAIN_RE = re.compile(r"^[a-z0-9](?:[a-z0-9-]*[a-z0-9])?(?:\.[a-z0-9](?:[a-z0-9-]*[a-z0-9])?)*$")


def norm(text) -> str:
    s = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).strip()
    return "" if s.lower() in EMPTY else s


def clean_domain(value) -> str:
    """'https://www.Esempio.it/' -> 'www.esempio.it'; '' se non sembra un nome di dominio."""
    s = cell(value).lower()
    s = re.sub(r"^[a-z]+://", "", s).split("/")[0].split("?")[0].strip().strip(".")
    return s if s and len(s) <= 253 and _DOMAIN_RE.match(s) else ""


def parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = cell(value)
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s[:10] if fmt == "%Y-%m-%d" else s, fmt).date()
        except ValueError:
            continue
    return None


def _map_header(row, fields) -> dict[int, str]:
    """Indice colonna -> campo, ogni campo una volta sola (vince la prima colonna)."""
    out, used = {}, set()
    for i, raw in enumerate(row):
        h = norm(raw)
        # un'intestazione e' corta e non e' un link o un percorso (quelli sono dati)
        if not h or len(h) > 50 or "://" in h or "<" in h or "→" in str(raw):
            continue
        for field, keys in fields:
            if field not in used and any(k in h for k in keys):
                out[i] = field
                used.add(field)
                break
    return out


def _find_header(rows, fields, required: str, min_cols: int = 2):
    for idx, row in enumerate(rows[:8]):
        cols = _map_header(row, fields)
        if required in cols.values() and len(cols) >= min_cols:
            return idx, cols
    return None, None


def _letter(i: int) -> str:
    out, n = "", i + 1
    while n:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out


def _sheet(rows: list[list], result: dict, sheet: str = "") -> str:
    """Classifica un foglio e ne aggiunge le righe a result. Ritorna il tipo riconosciuto."""
    key = norm(sheet)
    hidx, cols = _find_header(rows, DOMAIN_FIELDS, "domain")
    if cols is not None:
        fields = set(cols.values())
        inactive = "status_note" in fields and "mail_host" not in fields
        result["columns"][key] = ("domain", {_letter(i): f for i, f in cols.items()})
        target = result["inactive" if inactive else "domains"]
        for row in rows[hidx + 1:]:
            rec = {f: (row[i] if i < len(row) else None) for i, f in cols.items()}
            dom = clean_domain(rec.get("domain"))
            if not dom:
                if any(cell(v) for v in row):
                    result["skipped"].append(" ".join(cell(v) for v in row if cell(v))[:120])
                continue
            item = {"domain": dom, "active": not inactive, "_sheet": key}
            if "client" in fields:
                item["client"] = cell(rec.get("client"))[:190]
            for f in DOMAIN_TEXT:
                if f in fields:
                    item[f] = cell(rec.get(f))[:LIMITS.get(f, 4000)]
            if "hosting_expires" in fields:
                item["hosting_expires"] = parse_date(rec.get("hosting_expires"))
            target.append(item)
        return "inactive" if inactive else "domains"
    hidx, cols = _find_header(rows, PROVIDER_FIELDS, "name", min_cols=3)
    if cols is not None:
        fields = set(cols.values())
        result["columns"][key] = ("provider", {_letter(i): f for i, f in cols.items()})
        for row in rows[hidx + 1:]:
            filled = [v for v in (cell(x) for x in row) if v]
            if not filled:
                continue
            if len(filled) == 1 and len(filled[0]) >= 40:
                result["notes"].append((filled[0], key))      # consiglio generale (es. come spostare la posta)
                continue
            rec = {f: cell(row[i]) if i < len(row) else "" for i, f in cols.items()}
            pname = rec.get("name", "")[:190]
            if not pname:
                continue
            item = {"name": pname, "_sheet": key}
            for f in PROVIDER_TEXT:
                if f in fields:
                    item[f] = rec.get(f, "")
            result["providers"].append(item)
        return "providers"
    return ""


_REF = re.compile(r"""(?:\b(?:nella|nel|in|la|alla|della)\s+)?colonn[ae]\s+([A-Z]{1,2})\b"""
                  r"""(?:\s+(?:del|nel|dello)\s+foglio\s+(?:["“«']([^"”»']+)["”»']|([^\s.,;:)]+)))?""", re.I)


def rewrite_refs(text: str, result: dict, sheet: str = "") -> str:
    """"colonna G del foglio "Email per dominio"" -> dove sta quel dato nel pannello.
    Un riferimento che non si riesce a risolvere resta com'e'."""
    if not text or "colonn" not in text.lower():
        return text

    def sub(m):
        named = m.group(2) or m.group(3)
        target = norm(named) if named else sheet
        kind, cols = result["columns"].get(target, (None, {}))
        field = cols.get(m.group(1).upper())
        if kind == "domain" and field in FIELD_LABELS:
            return f'nel campo "{FIELD_LABELS[field]}" del dominio (Gestione domini → scheda)'
        if kind == "provider" and field in PROVIDER_LABELS:
            return f'nel campo "{PROVIDER_LABELS[field]}" del provider (Gestione domini → Provider)'
        return m.group(0)
    return _REF.sub(sub, text)


def parse_workbook(data: bytes, filename: str = "") -> dict:
    """Legge .xlsx o .csv. Ritorna {domains, inactive, providers, notes, sheets, skipped}."""
    result = {"domains": [], "inactive": [], "providers": [], "notes": [], "sheets": [], "skipped": [], "columns": {}}
    name = (filename or "").lower()
    if name.endswith(".csv") or (not data.startswith(b"PK") and not name.endswith((".xlsx", ".xlsm"))):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
        kind = _sheet(rows, result, filename or "CSV")
        result["sheets"].append({"name": filename or "CSV", "kind": kind})
    else:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as ex:  # noqa: BLE001
            raise ValueError(f"File Excel non leggibile: {ex}") from ex
        try:
            for ws in wb.worksheets:
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
                kind = _sheet(rows, result, ws.title)
                result["sheets"].append({"name": ws.title, "kind": kind})
        finally:
            wb.close()
    # un dominio sia attivo sia nell'elenco dei non attivi: vale quello attivo
    active = {d["domain"] for d in result["domains"]}
    result["inactive"] = [d for d in result["inactive"] if d["domain"] not in active]
    # testi che rimandano a colonne del foglio: diventano il posto del dato nel pannello
    for item in result["domains"] + result["inactive"]:
        for f in ("notes", "manage_howto", "status_note"):
            if item.get(f):
                item[f] = rewrite_refs(item[f], result, item["_sheet"])
    for item in result["providers"]:
        for f in PROVIDER_TEXT:
            if item.get(f):
                item[f] = rewrite_refs(item[f], result, item["_sheet"])
    notes = []
    for text, sheet in result["notes"]:
        text = rewrite_refs(text, result, sheet)
        if text not in notes:
            notes.append(text)
    result["notes"] = notes
    return result


# estensioni generiche piu' comuni (quelle di due lettere sono tutte nazionali e valgono sempre)
GTLDS = set("""com net org info biz name pro mobi asia tel travel jobs museum aero coop int edu gov mil
app dev page blog shop store online site website web tech cloud host hosting email digital agency studio design
media news live life world today company group solutions services consulting software systems network
center centre club events social art gallery photo photography pics music film video games game fun
xyz top win vip one ltd srl spa inc llc love guru expert academy school education training
hotel restaurant cafe bar pizza wine beer food farm garden house home immo realestate properties
health care clinic dental doctor fitness yoga beauty fashion style shoes bike auto car cars
casa moda bio eco green energy solar city tours travel holiday vacations
example test invalid localhost""".split())


def base_name(domain: str) -> str:
    """Nome che raggruppa i domini dello stesso sottocliente: esempio.com ed esempio.it,
    blog.esempio e shop.esempio.it, esempio.it e backend.esempio.it."""
    labels = [x for x in domain.lower().split(".") if x]
    if len(labels) < 2:
        return labels[0] if labels else domain
    last = labels[-1]
    if not ((len(last) == 2 and last.isalpha()) or last in GTLDS):
        return last                      # "blog.esempio": l'ultima parte non e' un'estensione
    if len(labels) >= 3 and labels[-2] in ("co", "com", "org", "net", "gov", "ac", "edu") and len(last) == 2:
        return labels[-3]                # esempio.co.uk
    return labels[-2]


def group_domains(items: list[dict]) -> dict[str, dict]:
    """Gruppi di domini per sottocliente: chiave -> {name, domains}. Vale la colonna "Cliente" del
    foglio se c'e', altrimenti il nome comune del dominio; il gruppo prende il nome del dominio
    principale (attivo, senza sottodominio, il piu' corto)."""
    groups: dict[str, dict] = {}
    for it in items:
        key = "c:" + norm(it["client"]) if it.get("client") else "d:" + base_name(it["domain"])
        g = groups.setdefault(key, {"name": it.get("client") or "", "domains": []})
        g["domains"].append(it)
    for g in groups.values():
        if not g["name"]:
            best = min(g["domains"], key=lambda d: (not d["active"], d["domain"].count("."), len(d["domain"]), d["domain"]))
            g["name"] = best["domain"]
    return groups


# ------------------------------------------------------------------ abbinamento ai provider
def provider_key(name: str) -> str:
    s = re.sub(r"\(.*?\)", " ", norm(name))
    return re.sub(r"[^a-z0-9.]+", " ", s).strip()


def _alternatives(name: str) -> list[str]:
    parts = [provider_key(p) for p in str(name or "").split("/")]
    full = provider_key(name)
    return [p for p in dict.fromkeys([full, *parts]) if p]


def _same(a: str, b: str) -> bool:
    return a == b or a.startswith(b + " ") or b.startswith(a + " ") or a.startswith(b + ".") or b.startswith(a + ".")


def match_provider(mail_host: str, providers) -> int | None:
    """Provider che corrisponde a "Email ospitate su": prima il nome identico, poi l'inizio
    del nome a parole intere ("Zoho Mail (EU)" -> "Zoho"). providers: oggetti con id e name."""
    hosts = _alternatives(mail_host)
    if not hosts:
        return None
    exact, loose = None, None
    for p in providers:
        alts = _alternatives(p.name)
        if any(h == a for h in hosts for a in alts):
            exact = exact or p.id
        elif any(_same(h, a) for h in hosts for a in alts):
            loose = loose or p.id
    return exact or loose


def folder_entry_name(path: str) -> str:
    """Nome della voce per una cartella: "Clienti" -> "Clienti", "Clienti/Agenzia" -> "Agenzia",
    "Clienti/Agenzia/Bar" -> "Agenzia / Bar" (le voci hanno un livello di sottoclienti solo)."""
    parts = [p for p in str(path or "").split("/") if p]
    return " / ".join(parts[1:]) if len(parts) > 1 else (parts[0] if parts else "")


def site_host(url: str) -> str:
    s = re.sub(r"^[a-z]+://", "", str(url or "").lower()).split("/")[0].split(":")[0]
    return s[4:] if s.startswith("www.") else s
